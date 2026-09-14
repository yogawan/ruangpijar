from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUT = Path("output/RuangPijar_MindCraft_2026.pptx")

# UI-derived palette from app/globals.css. The deck deliberately uses only
# native PowerPoint shapes and text: no image, logo, icon, or external asset.
BG = "FAF8F4"
SURFACE = "FFFFFF"
MUTED = "EEE5EC"
PEACH = "F8E7DF"
SAGE = "8FA58D"
SAGE_SOFT = "E7EFE5"
TEXT = "292525"
TEXT_MUTED = "746E6A"
BRAND = "5B3A52"
BRAND_DARK = "432B3D"
ACCENT = "E9A68D"
BORDER = "E8E2DC"
VIOLET = "6D28D9"
RED = "C8605A"

W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def rgb(value):
    return RGBColor.from_string(value)


def shape(slide, kind, x, y, w, h, fill, line=None, radius=False):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else kind,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.color.rgb = rgb(line or fill)
    return shp


def text(slide, value, x, y, w, h, size=18, color=TEXT, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         italic=False, margin=0.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return box


def line(slide, x1, y1, x2, y2, color=BORDER, width=1.0, dash=None):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = dash
    return ln


def base(slide, n, section, dark=False):
    bg = BRAND if dark else BG
    shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H, bg)
    # Native abstract glows / dots carry the app's soft ambient style.
    if not dark:
        glow = shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 11.7, -0.75, 2.25, 2.25, PEACH)
        glow.fill.transparency = 35
        glow.line.fill.background()
        glow2 = shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, -0.8, 6.55, 1.7, 1.7, SAGE_SOFT)
        glow2.fill.transparency = 18
        glow2.line.fill.background()
    text(slide, "RUANGPIJAR", 0.55, 0.30, 1.75, 0.25, 9, ACCENT if dark else BRAND, True, "Aptos Display")
    text(slide, section.upper(), 2.55, 0.30, 4.3, 0.25, 8.5, "E9D9E4" if dark else TEXT_MUTED, True)
    text(slide, f"{n:02d}", 12.17, 0.28, 0.55, 0.3, 10, "E9D9E4" if dark else TEXT_MUTED, True, align=PP_ALIGN.RIGHT)
    line(slide, 0.55, 0.73, 12.78, 0.73, "8B607D" if dark else BORDER, 0.65)


def title(slide, kicker, heading, sub=None, dark=False):
    text(slide, kicker.upper(), 0.7, 1.05, 5.6, 0.28, 10, ACCENT if dark else BRAND, True)
    text(slide, heading, 0.7, 1.38, 8.2, 1.05, 29, SURFACE if dark else TEXT, True, "Aptos Display")
    if sub:
        text(slide, sub, 0.72, 2.5, 7.6, 0.72, 13, "F4EAF0" if dark else TEXT_MUTED)


def pill(slide, label, x, y, w=None, fill=MUTED, fg=BRAND):
    w = w or max(0.8, len(label) * 0.085 + 0.35)
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 0.34, fill, fill, True)
    text(slide, label, x, y + 0.04, w, 0.2, 8, fg, True, align=PP_ALIGN.CENTER)


def card(slide, x, y, w, h, heading, body, num=None, fill=SURFACE, accent=BRAND):
    shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, BORDER, True)
    if num:
        text(slide, num, x + 0.24, y + 0.22, 0.55, 0.3, 12, accent, True, "Aptos Display")
    text(slide, heading, x + 0.24, y + (0.62 if num else 0.28), w - 0.48, 0.38, 14, TEXT, True)
    text(slide, body, x + 0.24, y + (1.08 if num else 0.76), w - 0.48, h - (1.25 if num else 0.96), 10.2, TEXT_MUTED)


def bullet(slide, x, y, lead, body, width=5.2, dark=False):
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y + 0.12, 0.12, 0.12, ACCENT, ACCENT)
    text(slide, lead, x + 0.25, y, width - 0.25, 0.25, 11.5, SURFACE if dark else TEXT, True)
    text(slide, body, x + 0.25, y + 0.29, width - 0.25, 0.43, 9.7, "F1E6ED" if dark else TEXT_MUTED)


def mini_face(slide, x, y, mood, fill):
    # A small, intentionally generic native-vector mood marker.
    shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, 0.55, 0.55, fill, fill)
    text(slide, mood, x, y + 0.08, 0.55, 0.28, 14, BRAND, True, align=PP_ALIGN.CENTER)


# 01 Cover
s = prs.slides.add_slide(blank)
base(s, 1, "MindCraft Web Competition 2026", True)
shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, 8.65, 0.9, 4.1, 4.1, BRAND_DARK, BRAND_DARK)
shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, 9.55, 1.8, 2.3, 2.3, ACCENT, ACCENT)
shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, 10.15, 2.4, 1.1, 1.1, SAGE_SOFT, SAGE_SOFT)
text(s, "RUANG KECIL UNTUK\nMEMAHAMI DIRIMU.", 0.75, 1.42, 6.9, 0.34, 11, ACCENT, True)
text(s, "RuangPijar", 0.72, 1.92, 7.1, 0.78, 42, SURFACE, True, "Aptos Display")
text(s, "Kamu tidak harus selalu\nbaik-baik saja.", 0.75, 2.86, 6.9, 1.1, 27, SURFACE, True, "Aptos Display")
text(s, "Aplikasi web untuk mencatat keseharian, mengenali pola, dan mencoba langkah kecil.", 0.76, 4.32, 5.7, 0.55, 13, "F1E6ED")
pill(s, "MINDCRAFT WEB COMPETITION 2026", 0.75, 5.3, 2.7, "8B607D", SURFACE)
text(s, "Building Digital Solutions, for Mental Well-being", 0.78, 5.84, 5.6, 0.3, 10.5, "F1E6ED")
text(s, "[NAMA TIM & ANGGOTA]", 0.78, 6.65, 4.3, 0.25, 9.5, ACCENT, True)

# 02 Problem
s = prs.slides.add_slide(blank); base(s, 2, "Pendahuluan & urgensi")
title(s, "Beban yang sering tak disadari", "Hari yang terlihat biasa,\nbelum tentu terasa biasa.", "Tuntutan akademik, sosial, dan pribadi dapat datang sekaligus.")
card(s, 0.7, 3.55, 3.75, 2.22, "Akademik", "Deadline, tugas, ujian, dan tuntutan untuk terus berkembang.", "01", PEACH, ACCENT)
card(s, 4.78, 3.55, 3.75, 2.22, "Keseharian", "Tidur, energi, rutinitas, dan waktu untuk beristirahat.", "02", SAGE_SOFT, SAGE)
card(s, 8.86, 3.55, 3.75, 2.22, "Sosial & personal", "Hubungan, keluarga, pekerjaan, dan hal-hal yang sering disimpan sendiri.", "03", MUTED, BRAND)
text(s, "Yang dibutuhkan bukan sekadar layanan tambahan, melainkan media pendukung yang mudah diakses, informatif, dan sesuai kebutuhan pengguna.", 0.72, 6.28, 11.4, 0.44, 13, BRAND, True)

# 03 Research
s = prs.slides.add_slide(blank); base(s, 3, "Pendahuluan & urgensi")
title(s, "Peluang solusi digital", "Teknologi bisa menjembatani\nakses yang masih terbatas.")
card(s, 0.7, 3.0, 5.85, 2.62, "Lattie dkk. (2019)", "Tinjauan sistematis atas 89 penelitian: 80% intervensi kesehatan mental digital untuk mahasiswa disampaikan lewat website; banyak yang efektif atau sebagian efektif bagi depresi, kecemasan, dan psychological well-being.", "01", SURFACE, BRAND)
card(s, 6.78, 3.0, 5.85, 2.62, "Naslund dkk. (2017)", "Teknologi digital dapat menjembatani akses saat layanan kesehatan mental konvensional terbatas.", "02", SURFACE, BRAND)
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 0.7, 6.03, 11.93, 0.62, SAGE_SOFT, SAGE_SOFT, True)
text(s, "Efektivitas tidak cukup diukur dari fitur: usability, penerimaan pengguna, pengalaman, keamanan, dan privasi sama pentingnya.", 0.98, 6.22, 11.3, 0.22, 10.7, TEXT, True, align=PP_ALIGN.CENTER)

# 04 Solution
s = prs.slides.add_slide(blank); base(s, 4, "Solusi")
title(s, "RuangPijar", "Catat sebentar. Pahami pola.\nCoba langkah kecil.", "Alat bantu refleksi pribadi untuk mengenali kondisi sehari-hari.")
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 8.75, 1.45, 3.9, 4.8, SURFACE, BORDER, True)
text(s, "CHECK-IN HARI INI", 9.1, 1.82, 2.4, 0.2, 8.8, TEXT_MUTED, True)
text(s, "Bagaimana perasaanmu?", 9.1, 2.35, 2.8, 0.3, 14, TEXT, True)
for i, (m, c) in enumerate([(":(", PEACH), (":|", MUTED), (":)", SAGE_SOFT), (":D", "F7DBCF")]): mini_face(s, 9.12 + i*0.78, 2.92, m, c)
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.1, 4.15, 3.18, 1.28, BRAND, BRAND, True)
text(s, "Ada ruang untuk\nsetiap harimu.", 9.38, 4.5, 2.62, 0.55, 16, SURFACE, True, "Aptos Display")
bullet(s, 0.75, 3.45, "Bukan aplikasi diagnosis", "RuangPijar tidak menggantikan bantuan psikolog atau konselor.")
bullet(s, 0.75, 4.48, "Dibangun untuk refleksi", "Suasana hati, energi, stres, dan hal-hal yang memengaruhinya dicatat dari keseharian pengguna.")
bullet(s, 0.75, 5.51, "Tidak harus selalu baik-baik saja", "Ruang untuk berhenti sejenak dan melihat diri dengan lebih jujur.")

# 05 Loop
s = prs.slides.add_slide(blank); base(s, 5, "Solusi")
title(s, "Satu siklus, empat langkah", "Bukan empat fitur terpisah.", "Setiap langkah memberi konteks bagi langkah berikutnya.")
steps = [("01", "Check-in", "Catat kondisi hari ini", BRAND, 0.8), ("02", "Jejak", "Lihat kembali perjalanan", SAGE, 3.88), ("03", "Insight", "Pahami pola yang muncul", ACCENT, 6.96), ("04", "Ruang", "Coba langkah kecil", BRAND, 10.04)]
for n, name, desc, col, x in steps:
    shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, x, 3.45, 1.55, 1.55, col, col)
    text(s, n, x, 3.7, 1.55, 0.24, 9, SURFACE, True, align=PP_ALIGN.CENTER)
    text(s, name, x, 4.02, 1.55, 0.26, 14, SURFACE, True, align=PP_ALIGN.CENTER)
    text(s, desc, x - 0.4, 5.35, 2.35, 0.45, 9.5, TEXT_MUTED, align=PP_ALIGN.CENTER)
for x in [2.43, 5.51, 8.59]:
    line(s, x, 4.22, x + 1.15, 4.22, BRAND, 1.8)
text(s, "Rutin check-in → data makin kaya → pola dan saran makin relevan → kembali check-in.", 1.2, 6.38, 10.95, 0.28, 12, BRAND, True, align=PP_ALIGN.CENTER)

# 06 Onboarding
s = prs.slides.add_slide(blank); base(s, 6, "Showcase produk")
title(s, "Mulai dengan kenalan singkat", "Akun yang menyambut, bukan membebani.")
card(s, 0.7, 3.0, 3.68, 2.7, "01 — Masuk", "Daftar lewat form manual atau satu klik Google. Email yang sama dikenali sebagai satu akun.", fill=SURFACE)
card(s, 4.83, 3.0, 3.68, 2.7, "02 — Kenalan", "Pilih fokus: mood, stres, energi, tidur, akademik, sosial, relasi, atau diri sendiri.", fill=SURFACE)
card(s, 8.96, 3.0, 3.68, 2.7, "03 — Atur ritme", "Tentukan frekuensi check-in dan jam pengingat yang paling nyaman. Semua bisa diubah di Profil.", fill=SURFACE)
text(s, "PREFERENSI AWAL", 0.73, 6.25, 1.8, 0.22, 9, BRAND, True)
pill(s, "Mood", 2.66, 6.17, 0.75, PEACH); pill(s, "Energi", 3.56, 6.17, 0.83, SAGE_SOFT, BRAND); pill(s, "Tidur", 4.54, 6.17, 0.75, MUTED); pill(s, "20:00", 5.44, 6.17, 0.8, BRAND, SURFACE)

# 07 Check-in
s = prs.slides.add_slide(blank); base(s, 7, "Showcase produk")
title(s, "Check-in", "Dua menit untuk berhenti sejenak.", "Lima pertanyaan santai; tidak ada jawaban salah dan sebagian besar boleh dilewati.")
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 7.75, 1.6, 4.85, 4.95, SURFACE, BORDER, True)
text(s, "CHECK-IN HARI INI", 8.1, 1.96, 2, 0.2, 8.5, TEXT_MUTED, True)
text(s, "Bagaimana mood-mu?", 8.1, 2.42, 3.0, 0.3, 15, TEXT, True)
for i, (m, c) in enumerate([(":(", PEACH), (":|", MUTED), (":)", SAGE_SOFT), (":D", "F7DBCF")]): mini_face(s, 8.1 + i*0.92, 2.98, m, c)
text(s, "Energi", 8.1, 3.93, 1.1, 0.2, 10, TEXT, True)
line(s, 8.1, 4.35, 11.95, 4.35, BORDER, 4); line(s, 8.1, 4.35, 10.72, 4.35, SAGE, 4)
shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, 10.62, 4.17, 0.35, 0.35, SAGE, SAGE)
text(s, "Stres", 8.1, 4.75, 1.1, 0.2, 10, TEXT, True)
line(s, 8.1, 5.17, 11.95, 5.17, BORDER, 4); line(s, 8.1, 5.17, 9.55, 5.17, ACCENT, 4)
shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, 9.45, 4.99, 0.35, 0.35, ACCENT, ACCENT)
bullet(s, 0.75, 3.35, "Data yang dicatat", "Mood, energi, stres; tidur dan beban akademik/sosial opsional; faktor dan cerita bebas.")
bullet(s, 0.75, 4.6, "Konsistensi yang diapresiasi", "Streak harian dan pengingat email hadir sebelum rangkaian check-in terputus.")

# 08 Jejak
s = prs.slides.add_slide(blank); base(s, 8, "Showcase produk")
title(s, "Jejak", "Melihat kembali perjalanan\nyang sudah dilewati.", "Kalender dan riwayat aktivitas menjadikan catatan kecil lebih mudah ditengok.")
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 7.15, 1.5, 5.48, 5.45, SURFACE, BORDER, True)
text(s, "SEPTEMBER 2026", 7.53, 1.91, 2.2, 0.25, 10, TEXT, True)
for i, d in enumerate(["S", "S", "R", "K", "J", "S", "M"]): text(s, d, 7.55 + i*.66, 2.33, .3, .18, 8, TEXT_MUTED, True, align=PP_ALIGN.CENTER)
for row in range(4):
    for col in range(7):
        x, y = 7.55 + col*.66, 2.72 + row*.62
        shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, .38, .38, SURFACE, BORDER)
        text(s, str(row*7+col+1), x, y+.1, .38, .12, 6.8, TEXT_MUTED, align=PP_ALIGN.CENTER)
for pos, face, c in [(2, ":)", SAGE_SOFT), (7, ":|", MUTED), (11, ":(", PEACH), (19, ":D", "F7DBCF")]:
    col, row = pos % 7, pos // 7
    mini_face(s, 7.47 + col*.66, 2.64 + row*.62, face, c)
line(s, 7.55, 5.47, 12.18, 5.47, BORDER, 1)
text(s, "RIWAYAT TERBARU", 7.55, 5.68, 1.7, .18, 8, TEXT_MUTED, True)
text(s, "Insight · Tidur berkaitan dengan stres", 7.55, 6.05, 3.9, .18, 9.2, TEXT, True)
bullet(s, .75, 3.25, "Kalender bulanan", "Tanggal yang terisi ditandai mood dan dapat dibuka untuk membaca ulang.")
bullet(s, .75, 4.55, "Satu rangkaian perjalanan", "Latihan dari Ruang dan pola dari Insight tersusun terbaru dulu dalam satu feed.")

# 09 Insight
s = prs.slides.add_slide(blank); base(s, 9, "Showcase produk")
title(s, "Insight", "Pola dari catatan sendiri,\ndibaca dengan jujur.", "Dibuat manual setelah data cukup; tingkat keyakinan menunjukkan sinyal, bukan kepastian.")
card(s, .7, 3.25, 3.76, 2.3, "Tren mood menurun", "Rata-rata mood berubah dalam jendela 14 hari terakhir.", "TREND", SURFACE, BRAND)
card(s, 4.79, 3.25, 3.76, 2.3, "Tidur ↔ stres", "Keterkaitan dihitung dari data check-in yang memiliki jam tidur.", "KORELASI", SURFACE, VIOLET)
card(s, 8.88, 3.25, 3.76, 2.3, "Faktor berulang", "Faktor yang sering muncul dapat ditandai sebagai pola.", "POLA", SURFACE, SAGE)
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, .7, 6.05, 11.94, .6, MUTED, MUTED, True)
text(s, "Insight = statistik sederhana dari data pengguna sendiri. Bukan AI, bukan diagnosis, dan “belum cukup data” adalah jawaban yang valid.", .98, 6.22, 11.35, .2, 10.4, BRAND, True, align=PP_ALIGN.CENTER)

# 10 Ruang
s = prs.slides.add_slide(blank); base(s, 10, "Showcase produk")
title(s, "Ruang", "Langkah kecil, tanpa tekanan.", "Latihan disarankan dari faktor yang paling sering muncul di catatan terbaru.")
for x, hd, bd, tag, col in [(.7, "Bernapas sejenak", "Ambil jeda pendek untuk menata napas.", "Napas · 3 menit", PEACH), (4.78, "Jurnal refleksi", "Tulis apa yang sedang memenuhi pikiran.", "Refleksi · 5 menit", SAGE_SOFT), (8.86, "Susun ulang prioritas", "Pilih satu hal kecil yang bisa dilakukan hari ini.", "Prioritas · 5 menit", MUTED)]:
    shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 3.0, 3.75, 2.75, SURFACE, BORDER, True)
    pill(s, tag, x+.25, 3.3, 1.55, col, BRAND)
    text(s, hd, x+.25, 3.95, 3.22, .28, 14, TEXT, True)
    text(s, bd, x+.25, 4.42, 3.15, .4, 10, TEXT_MUTED)
    shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x+.25, 5.07, 1.0, .34, BRAND, BRAND, True)
    text(s, "Mulai", x+.25, 5.15, 1.0, .15, 8.5, SURFACE, True, align=PP_ALIGN.CENTER)
text(s, "Mulai  →  Selesai atau Lewati", .75, 6.27, 11.8, .25, 13, BRAND, True, align=PP_ALIGN.CENTER)
text(s, "Keduanya wajar. Tidak ada penalti saat pengguna belum siap melanjutkan.", .75, 6.57, 11.8, .2, 9.8, TEXT_MUTED, align=PP_ALIGN.CENTER)

# 11 loop revisited
s = prs.slides.add_slide(blank); base(s, 11, "Showcase produk", True)
title(s, "Satu perjalanan yang menutup sendiri", "Dari mencatat,\nmenjadi memahami.", "Setiap pengalaman pengguna memperkaya siklus berikutnya.", True)
for i, (name, verb, x, col) in enumerate([("Check-in", "Catat", .9, ACCENT), ("Jejak", "Ingat", 3.92, SAGE), ("Insight", "Pahami", 6.94, "C389B0"), ("Ruang", "Bertindak", 9.96, ACCENT)]):
    shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 4.15, 2.38, 1.05, BRAND_DARK, "8B607D", True)
    text(s, f"{i+1:02d}", x+.22, 4.42, .4, .18, 8.5, col, True)
    text(s, name, x+.22, 4.35, 1.85, .24, 13, SURFACE, True)
    text(s, verb, x+.22, 4.7, 1.85, .18, 9.3, "E9D9E4")
for x in [3.28, 6.3, 9.32]: line(s, x, 4.68, x+.62, 4.68, ACCENT, 1.4)
text(s, "Bukan 4 fitur yang berdiri sendiri—melainkan satu kebiasaan refleksi yang saling menghidupkan.", .9, 6.15, 11.5, .33, 13, ACCENT, True, align=PP_ALIGN.CENTER)

# 12 supporting
s = prs.slides.add_slide(blank); base(s, 12, "Showcase produk")
title(s, "Fitur pendukung", "Agar kebiasaan kecil\nlebih mudah dimulai.")
card(s, .7, 3.1, 3.75, 2.58, "Panduan di tiap halaman", "Pengguna baru mendapat panduan singkat bergambar; dapat diputar ulang dari Profil.", "01", SAGE_SOFT, SAGE)
card(s, 4.79, 3.1, 3.75, 2.58, "Profil yang dapat diatur", "Ubah nama, foto, preferensi onboarding, dan keluar akun dari satu tempat.", "02", PEACH, ACCENT)
card(s, 8.88, 3.1, 3.75, 2.58, "Pengingat otomatis", "Email berjalan harian untuk pengguna yang streak check-in-nya akan putus.", "03", MUTED, BRAND)

# 13 Architecture
s = prs.slides.add_slide(blank); base(s, 13, "Dapur pacu & proses")
title(s, "Arsitektur teknologi", "Satu aplikasi web,\nlapisannya saling terhubung.", "Implementasi saat ini memakai Next.js 16, React 19, MongoDB, dan NextAuth.")
layers = [("Pengguna", "Browser", BRAND), ("Aplikasi", "Next.js App Router + React + TypeScript", SAGE), ("API", "Route handlers · sesi JWT · validasi", ACCENT), ("Data & layanan", "MongoDB/Mongoose · Gmail SMTP · Cloudinary", "C389B0")]
for i, (a,b,c) in enumerate(layers):
    y = 3.0 + i*.78
    shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.1, y, 11.1, .55, SURFACE, BORDER, True)
    shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.1, y, 2.0, .55, c, c, True)
    text(s, a, 1.28, y+.16, 1.65, .18, 10, SURFACE, True)
    text(s, b, 3.38, y+.16, 8.2, .18, 10, TEXT, True)
    if i < 3: line(s, 6.65, y+.56, 6.65, y+.76, BRAND, 1.2)
text(s, "Autentikasi: NextAuth v5 (JWT) · UI: Tailwind v4 · Animasi: GSAP · Onboarding: react-joyride · Pengingat: Vercel Cron + Nodemailer", .82, 6.45, 11.72, .3, 9.8, TEXT_MUTED, align=PP_ALIGN.CENTER)

# 14 Insight technical
s = prs.slides.add_slide(blank); base(s, 14, "Dapur pacu & proses")
title(s, "Cara kerja Insight", "Statistik sederhana,\naturan yang transparan.", "Tidak ada model AI/LLM: proses berjalan dari check-in pengguna dalam 14 hari terakhir.")
for x, hd, bd, col in [(.7, "01 · Ambil data", "Minimal 3 check-in dalam jendela 14 hari.", SAGE), (4.78, "02 · Hitung sinyal", "Tren, korelasi Pearson, dan faktor yang berulang.", ACCENT), (8.86, "03 · Tampilkan", "Simpan insight dan tingkat keyakinannya.", BRAND)]:
    shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, 3.05, 3.75, 1.28, SURFACE, BORDER, True)
    shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, x+.25, 3.38, .45, .45, col, col)
    text(s, hd, x+.86, 3.28, 2.45, .2, 10.5, TEXT, True)
    text(s, bd, x+.86, 3.63, 2.42, .3, 8.6, TEXT_MUTED)
for x in [4.45, 8.53]: line(s, x, 3.69, x+.33, 3.69, BRAND, 1.4)
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, .7, 4.85, 11.94, 1.18, MUTED, MUTED, True)
text(s, "Aturan yang digunakan", .98, 5.12, 2.4, .2, 10.5, BRAND, True)
text(s, "Tren: perubahan rata-rata ≥ 0,75  ·  Korelasi tidur–stres: |r| ≥ 0,40  ·  Faktor: muncul pada ≥ 40% check-in", .98, 5.48, 10.9, .22, 10, TEXT, True)
text(s, "Keyakinan dibatasi 0–100%; jika data belum cukup, sistem mengembalikan “Belum cukup data untuk membuat insight baru.”", .7, 6.45, 11.9, .23, 10.5, TEXT_MUTED, align=PP_ALIGN.CENTER)

# 15 Roadmap
s = prs.slides.add_slide(blank); base(s, 15, "Visi, eksekusi & penutup")
title(s, "Arah pengembangan", "Membuat dukungan terasa\nmakin relevan dan mandiri.", "Roadmap berikut dibaca sebagai peluang pertumbuhan, bukan kekurangan produk.")
for i, (hd, bd, x, c) in enumerate([("Personalisasi lebih dalam", "Preferensi onboarding terhubung ke saran Ruang dan pola Insight.", .7, PEACH), ("Kemandirian akun", "Reset password, verifikasi email, dan opsi hapus akun mandiri.", 4.78, SAGE_SOFT), ("Insight proaktif", "Dorongan berkala, tidak hanya menunggu tombol “Cari pola baru”.", 8.86, MUTED)]):
    card(s, x, 3.35, 3.75, 2.35, hd, bd, f"0{i+1}", c, BRAND)

# 16 values
s = prs.slides.add_slide(blank); base(s, 16, "Visi, eksekusi & penutup", True)
title(s, "Nilai & batasan yang dijaga", "Transparan, privat,\ndan tidak menghakimi.", "Batasan ini adalah bagian dari desain, bukan catatan kaki.", True)
bullet(s, .82, 3.35, "Bukan diagnosis", "RuangPijar bukan pengganti bantuan profesional seperti psikolog atau konselor.", 5.1, True)
bullet(s, .82, 4.48, "Privasi milik pengguna", "Setiap pengguna hanya dapat melihat catatan miliknya sendiri.", 5.1, True)
bullet(s, 6.9, 3.35, "Statistik, bukan black box", "Insight dihitung transparan dari data check-in pengguna sendiri—bukan AI.", 5.1, True)
bullet(s, 6.9, 4.48, "Tanpa AI atau Web3/blockchain", "Tidak ada klaim teknologi yang tidak digunakan oleh produk.", 5.1, True)

# 17 CTA
s = prs.slides.add_slide(blank); base(s, 17, "Visi, eksekusi & penutup")
title(s, "Mari coba alurnya", "Dua menit untuk memulai\nrefleksi yang lebih sadar.")
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, .72, 3.25, 6.2, 2.5, SURFACE, BORDER, True)
text(s, "ALUR LIVE DEMO", 1.06, 3.58, 1.8, .18, 9, BRAND, True)
text(s, "Daftar  →  Kenalan singkat  →  Check-in  →\nJejak  →  Cari pola baru  →  Ruang", 1.06, 4.08, 5.32, .55, 14, TEXT, True, "Aptos Display")
text(s, "[URL DEMO]", 1.06, 5.1, 2.2, .22, 10.5, TEXT_MUTED, True)
shape(s, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 8.42, 2.5, 3.15, 3.15, BRAND, BRAND, True)
text(s, "QR\nDEMO", 8.42, 3.35, 3.15, .55, 24, SURFACE, True, "Aptos Display", PP_ALIGN.CENTER)
text(s, "Ganti kotak ini dengan QR code demo saat tersedia.", 8.0, 6.03, 4.0, .28, 9.4, TEXT_MUTED, align=PP_ALIGN.CENTER)

# 18 Close
s = prs.slides.add_slide(blank); base(s, 18, "Visi, eksekusi & penutup", True)
shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, 8.8, 1.0, 3.45, 3.45, BRAND_DARK, BRAND_DARK)
shape(s, MSO_AUTO_SHAPE_TYPE.OVAL, 9.62, 1.82, 1.82, 1.82, SAGE_SOFT, SAGE_SOFT)
text(s, "RUANGPIJAR", .78, 1.3, 2.1, .25, 10, ACCENT, True)
text(s, "Kamu tidak harus\nselalu baik-baik saja.", .75, 1.85, 7.0, 1.15, 32, SURFACE, True, "Aptos Display")
text(s, "Kontribusi kecil mahasiswa Informatika untuk membantu keseharian terasa lebih dipahami.", .78, 3.4, 5.85, .48, 13, "F1E6ED")
line(s, .78, 4.55, 5.9, 4.55, "8B607D", 1)
text(s, "Terima kasih.", .78, 5.02, 3.2, .4, 19, ACCENT, True, "Aptos Display")
text(s, "MindCraft Web Competition 2026", .8, 6.35, 3.2, .22, 9.5, "E9D9E4")

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(OUT)
